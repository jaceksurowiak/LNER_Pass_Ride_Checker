"""
LNER (TRENT) - Pass Ride Checker
Checks other operators' trains (Brand blank) in a TRACS pass-trips export
against Realtime Trains (RTT).

VERSION FOR THE OLD RTT API (secure.realtimetrains.co.uk/api) – works until 31 March 2027.
From then on use app2.py, which uses the new token-based API.

Run:  streamlit run app.py
Secrets:  RTT_USER = "...", RTT_PASS = "..."
"""
import io
import re
import time
import threading
import datetime as dt
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Optional, Dict, Set, List, Tuple, Callable

import pandas as pd
import requests
import streamlit as st

try:   # lets the app remember settings in your browser
    from streamlit_js_eval import streamlit_js_eval
except ImportError:
    streamlit_js_eval = None
from requests.auth import HTTPBasicAuth

RTT_BASE = "https://secure.realtimetrains.co.uk/api"
DEFAULT_RAILREFS_PATH = "RailReferences.csv"
# Parallel RTT requests. Change DEFAULT_WORKERS to set the default, or add
# RTT_WORKERS = 10 to Streamlit Secrets to override it without editing code.
DEFAULT_WORKERS = 8
MAX_WORKERS = 16

STATUS_OK = "OK"
STATUS_FAIL = "FAIL"
STATUS_NC = "NOT CHECKED"

TRIP_KEY = ["date", "headcode", "origin_crs", "dest_crs", "exp_dep", "exp_arr"]


# ============================================================
# Text / time helpers
# ============================================================
def key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s).strip().lower())


def norm_name(s: str) -> str:
    s = str(s).strip().lower()
    s = re.sub(r"[^a-z0-9 ]+", "", s)
    s = re.sub(r"\s+", " ", s)
    return s


def find_col(df: pd.DataFrame, *candidates: str, allow_contains: bool = True) -> Optional[str]:
    cols = list(df.columns)
    km = {key(c): c for c in cols}
    for w in candidates:
        if key(w) in km:
            return km[key(w)]
    if not allow_contains:
        return None
    for w in candidates:
        kw = key(w)
        for c in cols:
            if kw and kw in key(c):
                return c
    return None


def parse_pass_hhmm(x) -> Optional[str]:
    """'05:32' or '05:32h' (half-minute working time) -> '05:32'."""
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return None
    s = str(x).strip().rstrip("hH")
    m = re.match(r"^(\d{1,2}):(\d{2})$", s)
    if not m:
        return None
    hh, mm = int(m.group(1)), int(m.group(2))
    if 0 <= hh <= 23 and 0 <= mm <= 59:
        return f"{hh:02d}:{mm:02d}"
    return None


def rtt_to_hhmm(x) -> Optional[str]:
    """RTT '0532' / '053230' -> '05:32'."""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    s = re.sub(r"\D", "", str(x))
    if len(s) in (4, 6):
        return f"{s[:2]}:{s[2:4]}"
    return None


def parse_date(x) -> Optional[dt.date]:
    if x is None or (not isinstance(x, str) and pd.isna(x)):
        return None
    s = str(x).strip()
    if not s:
        return None
    try:
        return pd.to_datetime(s, dayfirst=True, errors="raise").date()
    except Exception:
        return None


def in_window(day: dt.date, start, end) -> bool:
    if isinstance(start, dt.date) and day < start:
        return False
    if isinstance(end, dt.date) and day > end:
        return False
    return True


# ============================================================
# Rail day index: Sat=1, Sun=2, Mon=3, Tue=4, Wed=5, Thu=6, Fri=7
# ============================================================
RAIL_ORDER = ["Saturday", "Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
RAIL_BASE_MON_SAT = {1, 3, 4, 5, 6, 7}
RAIL_BASE_ALL_DAYS = {1, 2, 3, 4, 5, 6, 7}


def rail_day_index(day: dt.date) -> int:
    return ((day.weekday() + 2) % 7) + 1


def shift_rail_days(days: Set[int], n: int = 1) -> Set[int]:
    return {((d - 1 + n) % 7) + 1 for d in days}


def parse_days_run(value) -> Set[int]:
    """
    Days-run code -> rail-day indexes (Sat=1..Fri=7).
    Tokens: MO TO WO ThO FO SO Su.  Compact: M T W Th F S Su + trailing O (only) / X (except).
    X codes use a Mon–Sat base unless Sunday is mentioned (Sunday runs appear as Su rows).
    """
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return set()
    raw = str(value).strip()
    if not raw:
        return set()

    rail_map = {"SO": 1, "SU": 2, "MO": 3, "TO": 4, "WO": 5, "THO": 6, "FO": 7}
    days = set()
    for t in re.split(r"[,\s;/]+", raw):
        t_up = t.strip().upper().replace("TH0", "THO")
        if t_up in rail_map:
            days.add(rail_map[t_up])
    if days:
        return days

    blob = re.sub(r"[\s,;/]+", "", raw)
    mode = None
    if blob and blob[-1].upper() in ("O", "X"):
        mode = blob[-1].upper()
        blob = blob[:-1]
    base = RAIL_BASE_ALL_DAYS if "su" in blob.lower() else RAIL_BASE_MON_SAT

    picked: Set[int] = set()
    i = 0
    while i < len(blob):
        two = blob[i:i + 2].lower()
        one = blob[i:i + 1].upper()
        if two == "th":
            picked.add(6); i += 2; continue
        if two == "su":
            picked.add(2); i += 2; continue
        picked.update({"M": {3}, "T": {4}, "W": {5}, "F": {7}, "S": {1}}.get(one, set()))
        i += 1

    if mode == "O":
        return picked
    if mode == "X":
        return (base - picked) if picked else set(base)
    return picked


def journey_day_offset(jdays: Set[int], ddays: Set[int], diagram_date: dt.date) -> Optional[int]:
    """
    Which calendar day a trip runs for a given DIAGRAM date.
      0    -> same day as the diagram
      1    -> next day (after midnight): JourneyDays Run = DiagramDays Run shifted one day
      None -> trip doesn't run for this diagram date
    """
    rd_today = rail_day_index(diagram_date)
    rd_next = rail_day_index(diagram_date + dt.timedelta(days=1))
    if ddays and jdays:
        if jdays == ddays:
            return 0 if rd_today in jdays else None
        if jdays == shift_rail_days(ddays, 1):
            return 1 if rd_next in jdays else None
    if rd_today in jdays:
        return 0
    if rd_next in jdays:
        return 1
    return None


# ============================================================
# RailReferences
# ============================================================
def read_railrefs(src) -> pd.DataFrame:
    return pd.read_csv(src, header=None, names=["tiploc", "crs", "description"])


def build_desc_to_crs(rail_refs: pd.DataFrame) -> Dict[str, str]:
    d: Dict[str, str] = {}
    for desc, crs in zip(rail_refs["description"], rail_refs["crs"]):
        n = norm_name(desc)
        c = str(crs).strip().upper() if not pd.isna(crs) else ""
        if n and c and n not in d:
            d[n] = c
    for k, v in {
        "kings cross": "KGX", "london kings cross": "KGX", "kings x": "KGX",
        "edinburgh": "EDB", "edinburgh waverley": "EDB",
    }.items():
        d.setdefault(norm_name(k), v)
    return d


# ============================================================
# Pass-trips file
# ============================================================
COLS = {
    "brand": ("Brand",),
    "headcode": ("Headcode",),
    "origin": ("JourneyOrigin", "Journey Origin"),
    "dest": ("JourneyDestination", "Journey Destination"),
    "dep": ("JourneyDeparture", "Journey Departure"),
    "arr": ("JourneyArrival", "Journey Arrival"),
    "jdays": ("JourneyDays Run", "Journey Days Run", "JourneyDaysRun"),
    "train_start": ("TrainStart Date", "Train Start Date"),
    "train_end": ("TrainEnd Date", "Train End Date"),
    "diag_start": ("DiagramStart Date", "Diagram Start Date"),
    "diag_end": ("DiagramEnd Date", "Diagram End Date"),
    "resource": ("Resource",),
    "plan": ("DiagramPlan Type", "Diagram Plan Type"),
    "depot": ("DiagramDepot", "Diagram Depot"),
    "did": ("DiagramID", "Diagram Id", "Diagram ID"),
    "ddays": ("DiagramDays Run", "Diagram Days Run", "DiagramDaysRun"),
}
REQUIRED = ["brand", "headcode", "origin", "dest", "dep", "arr", "jdays"]


def load_pass_trips(src) -> pd.DataFrame:
    return pd.read_csv(src, skiprows=3, dtype=str, keep_default_na=False)


def map_columns(df: pd.DataFrame) -> Tuple[Dict[str, Optional[str]], List[str]]:
    cols = {}
    for k, names in COLS.items():
        exact_only = k in ("jdays", "ddays")
        cols[k] = find_col(df, *names, allow_contains=not exact_only)
    missing = [COLS[k][0] for k in REQUIRED if cols[k] is None]
    return cols, missing


def prepare_trips(df: pd.DataFrame, cols: Dict[str, Optional[str]], desc_to_crs: Dict[str, str]) -> pd.DataFrame:
    """Keep Brand-blank rows only and add parsed fields."""
    brand = df[cols["brand"]].astype(str).str.strip()
    t = df[brand == ""].copy()

    def col(k, default=""):
        return t[cols[k]].astype(str).str.strip() if cols.get(k) else pd.Series(default, index=t.index)

    out = pd.DataFrame(index=t.index)
    out["headcode"] = col("headcode")
    out["exp_from"] = col("origin")
    out["exp_to"] = col("dest")
    out["exp_dep"] = col("dep").apply(parse_pass_hhmm)
    out["exp_arr"] = col("arr").apply(parse_pass_hhmm)
    out["journey_days"] = col("jdays")
    out["jdays_set"] = out["journey_days"].apply(parse_days_run)
    out["diagram_days"] = col("ddays")
    out["ddays_set"] = out["diagram_days"].apply(parse_days_run)
    out["origin_crs"] = out["exp_from"].apply(lambda x: desc_to_crs.get(norm_name(x)))
    out["dest_crs"] = out["exp_to"].apply(lambda x: desc_to_crs.get(norm_name(x)))
    for k in ("train_start", "train_end", "diag_start", "diag_end"):
        out[k] = col(k).apply(parse_date)
    out["resource"] = col("resource")
    out["plan_type"] = col("plan")
    out["depot"] = col("depot")
    out["diagram_id"] = col("did")
    out["diagram"] = (out["depot"] + " " + out["diagram_id"] + " · " + out["plan_type"] + " " + out["diagram_days"]).str.strip()
    return out.reset_index(drop=True)


# ============================================================
# Railway weeks (Network Rail planning weeks)
# Weeks run Saturday–Friday. Week 1 starts on the Saturday NEAREST 1 April,
# so a rail year has 52 or 53 weeks. Checked against Network Rail's
# Calendar of Milestone dates: Week 53 = 28/03/2026, Week 1 = 04/04/2026,
# Week 38 = 19/12/2026.
# ============================================================
def rail_year_start(year: int) -> dt.date:
    anchor = dt.date(year, 4, 1)
    back = (anchor.weekday() - 5) % 7            # days back to the Saturday on or before 1 April
    return anchor - dt.timedelta(days=back) if back <= 3 else anchor + dt.timedelta(days=7 - back)


def rail_week_of(d: dt.date) -> Tuple[int, int]:
    """date -> (rail year it starts in, week number)."""
    start = rail_year_start(d.year)
    if d < start:
        start = rail_year_start(d.year - 1)
    return start.year, (d - start).days // 7 + 1


def rail_week_dates(rail_year: int, week: int) -> List[dt.date]:
    sat = rail_year_start(rail_year) + dt.timedelta(weeks=week - 1)
    return [sat + dt.timedelta(days=i) for i in range(7)]


def weeks_in_rail_year(rail_year: int) -> int:
    return (rail_year_start(rail_year + 1) - rail_year_start(rail_year)).days // 7


def detect_rail_week(trips: pd.DataFrame) -> Optional[Tuple[int, int]]:
    """Railway week the file is for: the week most diagram start dates fall in."""
    weeks = [rail_week_of(d) for d in trips["diag_start"] if isinstance(d, dt.date)]
    if not weeks:
        return None
    return max(set(weeks), key=weeks.count)


def build_expected(trips: pd.DataFrame, diagram_dates: List[dt.date]) -> pd.DataFrame:
    """One row per trip per chosen DIAGRAM date; 'date' is the day the train actually runs."""
    rows: List[dict] = []
    for day in sorted(diagram_dates):
        rd = rail_day_index(day)
        for r in trips.itertuples(index=False):
            ddays, jdays = r.ddays_set, r.jdays_set
            if ddays:
                if rd not in ddays or not in_window(day, r.diag_start, r.diag_end):
                    continue
                offset = journey_day_offset(jdays, ddays, day)
                if offset is None:
                    continue
            else:  # no DiagramDays Run in the file: journey day = calendar day
                if rd not in jdays or not in_window(day, r.diag_start, r.diag_end):
                    continue
                offset = 0
            run_day = day + dt.timedelta(days=offset)
            if not (in_window(day, r.train_start, r.train_end) or in_window(run_day, r.train_start, r.train_end)):
                continue
            rows.append({
                "diagram_date": day.isoformat(),
                "diagram_day": day.strftime("%A"),
                "rail_day": rd,
                "date": run_day.isoformat(),
                "after_midnight": offset == 1,
                "headcode": r.headcode,
                "exp_from": r.exp_from, "exp_to": r.exp_to,
                "exp_dep": r.exp_dep, "exp_arr": r.exp_arr,
                "origin_crs": r.origin_crs, "dest_crs": r.dest_crs,
                "resource": r.resource, "depot": r.depot, "diagram_id": r.diagram_id,
                "plan_type": r.plan_type, "diagram_days": r.diagram_days, "diagram": r.diagram,
            })
    return pd.DataFrame(rows)


def expected_for_days(trips: pd.DataFrame, week_dates: List[dt.date], days: List[str]) -> pd.DataFrame:
    """
    Trips for the chosen days of the week. A day includes
      - every trip on that day's diagrams, and
      - after-midnight trips that RUN on that day from the previous day's diagram,
        even when that diagram is outside the chosen days (e.g. Friday diagram,
        train at 00:17 on Saturday when only Saturday is picked).
    """
    exp = build_expected(trips, week_dates)
    if exp.empty:
        return exp
    picked = {DAY_SHORT.index(d) for d in days}               # 0 = Sat … 6 = Fri
    diag_idx = exp["rail_day"] - 1
    run_idx = exp["date"].map(lambda d: rail_day_index(dt.date.fromisoformat(d)) - 1)
    keep = diag_idx.isin(picked) | (exp["after_midnight"] & run_idx.isin(picked))
    out = exp[keep].copy()
    out["run_day"] = out["date"].map(lambda d: dt.date.fromisoformat(d).strftime("%A"))
    return out.reset_index(drop=True)


DAY_SHORT = ["Sat", "Sun", "Mon", "Tue", "Wed", "Thu", "Fri"]
WORKERS_STORE = "pass_ride_checker_parallel_requests"


def remembered_workers() -> Optional[int]:
    """The default saved in this browser, or None if there isn't one (yet)."""
    if streamlit_js_eval is None:
        return None
    val = streamlit_js_eval(js_expressions=f"localStorage.getItem('{WORKERS_STORE}')", key="read_saved_workers")
    try:
        return max(1, min(MAX_WORKERS, int(val)))
    except (TypeError, ValueError):
        return None


# ============================================================
# RTT clients
# ============================================================
class RTTClient:
    """Thread-safe RTT client with a shared session per thread and a small TTL cache."""

    def __init__(self, user: str, pw: str):
        self.auth = HTTPBasicAuth(user, pw)
        self._local = threading.local()
        self._cache: Dict[tuple, Tuple[float, dict]] = {}
        self._lock = threading.Lock()

    def _session(self) -> requests.Session:
        s = getattr(self._local, "s", None)
        if s is None:
            s = requests.Session()
            s.auth = self.auth
            adapter = requests.adapters.HTTPAdapter(pool_connections=4, pool_maxsize=16)
            s.mount("https://", adapter)
            self._local.s = s
        return s

    def _get(self, url: str, ttl: int) -> dict:
        now = time.time()
        with self._lock:
            hit = self._cache.get((url,))
            if hit and now - hit[0] < ttl:
                return hit[1]
        for attempt in range(3):
            r = self._session().get(url, timeout=30)
            if r.status_code == 429 and attempt < 2:   # throttled: back off and retry
                time.sleep(1.5 * (attempt + 1))
                continue
            r.raise_for_status()
            data = r.json()
            with self._lock:
                self._cache[(url,)] = (time.time(), data)
            return data
        raise RuntimeError("RTT rate limit")

    def location_search(self, crs: str, run_date: dt.date) -> dict:
        return self._get(f"{RTT_BASE}/json/search/{crs}/{run_date:%Y/%m/%d}", ttl=300)

    def service_detail(self, uid: str, run_date_iso: str) -> dict:
        y, m, d = str(run_date_iso).split("-")
        return self._get(f"{RTT_BASE}/json/service/{uid}/{y}/{m}/{d}", ttl=3600)


# ============================================================
# Checking engine
# ============================================================
def flatten_services(payload: dict) -> pd.DataFrame:
    rows = []
    for s in payload.get("services", []) or []:
        ld = s.get("locationDetail") or {}
        rows.append({
            "trainIdentity": (s.get("trainIdentity") or "").strip(),
            "serviceUid": s.get("serviceUid"),
            "runDate": s.get("runDate"),
            "operator": s.get("atocName") or "",
            "plannedCancel": bool(s.get("plannedCancel", False)),
            "dep": rtt_to_hhmm(ld.get("gbttBookedDeparture") or ld.get("wttBookedDeparture")),
        })
    return pd.DataFrame(rows, columns=["trainIdentity", "serviceUid", "runDate", "operator", "plannedCancel", "dep"])


CANCELLED = {"CANCELLED_CALL", "CANCELLED_PASS"}


def extract_segment(detail: dict, origin_crs: str, dest_crs: str) -> dict:
    seg = {"from_found": False, "to_found": False, "act_from": None, "act_to": None,
           "act_dep": None, "act_arr": None, "from_cancelled": False, "to_cancelled": False}
    for loc in detail.get("locations", []) or []:
        ld = loc.get("locationDetail") or loc
        crs = (loc.get("crs") or "").strip().upper()
        cancelled = (ld.get("displayAs") in CANCELLED) or bool(ld.get("cancelReasonCode"))
        if not seg["from_found"] and crs == origin_crs:
            seg.update(from_found=True, act_from=loc.get("description"), from_cancelled=cancelled,
                       act_dep=rtt_to_hhmm(ld.get("gbttBookedDeparture") or ld.get("wttBookedDeparture")))
        elif seg["from_found"] and crs == dest_crs:
            seg.update(to_found=True, act_to=loc.get("description"), to_cancelled=cancelled,
                       act_arr=rtt_to_hhmm(ld.get("gbttBookedArrival") or ld.get("wttBookedArrival")))
            break
    return seg


def _result(status: str, error: str = "", **kw) -> dict:
    base = {"status": status, "error": error, "operator": "", "planned_cancel": False,
            "serviceUid": "", "runDate": "", "act_from": None, "act_to": None,
            "act_dep": None, "act_arr": None, "origin_rtt_deps": ""}
    base.update(kw)
    return base


def pick_candidate(trip: dict, loc_df: pd.DataFrame):
    """Return (candidate_row, None) or (None, result_dict) for a trip using the origin search."""
    cand = loc_df[loc_df["trainIdentity"] == trip["headcode"]]
    if cand.empty:
        return None, _result(STATUS_FAIL, "Headcode not found at origin on this date")
    deps = sorted(set(cand["dep"].dropna()))
    extra = {"origin_rtt_deps": ", ".join(deps), "operator": ", ".join(sorted(set(cand["operator"])))}
    if trip["exp_dep"]:
        match = cand[cand["dep"] == trip["exp_dep"]]
        if match.empty:
            return None, _result(STATUS_FAIL, f"Departs origin at {', '.join(deps) or '—'} in RTT, not {trip['exp_dep']}",
                                 act_dep=deps[0] if deps else None, **extra)
        cand = match
    c = cand.iloc[0]
    if c["plannedCancel"]:
        return None, _result(STATUS_FAIL, "Planned cancellation in RTT", planned_cancel=True,
                             serviceUid=c["serviceUid"], runDate=c["runDate"], act_dep=c["dep"],
                             operator=c["operator"], origin_rtt_deps=extra["origin_rtt_deps"])
    return c, extra


def judge(trip: dict, c, extra: dict, detail: dict) -> dict:
    seg = extract_segment(detail, trip["origin_crs"], trip["dest_crs"])
    common = dict(operator=c["operator"], serviceUid=c["serviceUid"], runDate=c["runDate"],
                  act_from=seg["act_from"], act_to=seg["act_to"], act_dep=seg["act_dep"],
                  act_arr=seg["act_arr"], origin_rtt_deps=extra["origin_rtt_deps"])
    if not seg["from_found"]:
        return _result(STATUS_FAIL, f"Train doesn't call at {trip['exp_from']} ({trip['origin_crs']})", **common)
    if not seg["to_found"]:
        return _result(STATUS_FAIL, f"Train doesn't call at {trip['exp_to']} ({trip['dest_crs']}) after origin", **common)
    if seg["from_cancelled"]:
        return _result(STATUS_FAIL, f"Call at {trip['exp_from']} cancelled", **common)
    if seg["to_cancelled"]:
        return _result(STATUS_FAIL, f"Call at {trip['exp_to']} cancelled", **common)
    if trip["exp_dep"] and seg["act_dep"] and trip["exp_dep"] != seg["act_dep"]:
        return _result(STATUS_FAIL, f"Departure differs: RTT {seg['act_dep']}", **common)
    if trip["exp_arr"] and seg["act_arr"] and trip["exp_arr"] != seg["act_arr"]:
        return _result(STATUS_FAIL, f"Arrival differs: RTT {seg['act_arr']}", **common)
    return _result(STATUS_OK, "", **common)


def run_checks(expected: pd.DataFrame, client, workers: int = DEFAULT_WORKERS,
               progress: Optional[Callable[[int, int, int], None]] = None) -> pd.DataFrame:
    """Check every unique trip once (in parallel) and spread results back to every diagram row."""
    trips = expected.drop_duplicates(TRIP_KEY).to_dict("records")
    results: Dict[tuple, dict] = {}
    pending = []
    for t in trips:
        k = tuple(t[c] for c in TRIP_KEY)
        if not isinstance(t["origin_crs"], str):
            results[k] = _result(STATUS_NC, f"'{t['exp_from']}' not in RailReferences (no CRS)")
        elif not isinstance(t["dest_crs"], str):
            results[k] = _result(STATUS_NC, f"'{t['exp_to']}' not in RailReferences (no CRS)")
        else:
            pending.append((k, t))

    total = len(trips)
    done = len(results)
    fails = 0

    def tick(res):
        nonlocal done, fails
        done += 1
        fails += res["status"] == STATUS_FAIL
        if progress:
            progress(done, total, fails)

    # Phase 1: one station search per (origin, date), in parallel
    searches = sorted({(t["origin_crs"], t["date"]) for _, t in pending})
    loc_data: Dict[tuple, object] = {}
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(client.location_search, crs, dt.date.fromisoformat(d)): (crs, d) for crs, d in searches}
        for f in as_completed(futs):
            try:
                loc_data[futs[f]] = flatten_services(f.result())
            except Exception as e:
                loc_data[futs[f]] = e

    # Phase 2: pick the service at origin, then fetch its calling points, in parallel
    with ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {}
        for k, t in pending:
            loc = loc_data[(t["origin_crs"], t["date"])]
            if isinstance(loc, Exception):
                results[k] = _result(STATUS_FAIL, f"RTT station search failed: {loc}"); tick(results[k]); continue
            c, info = pick_candidate(t, loc)
            if c is None:
                results[k] = info; tick(info); continue
            futs[ex.submit(client.service_detail, c["serviceUid"], c["runDate"])] = (k, t, c, info)
        for f in as_completed(futs):
            k, t, c, info = futs[f]
            try:
                results[k] = judge(t, c, info, f.result())
            except Exception as e:
                results[k] = _result(STATUS_FAIL, f"RTT service lookup failed: {e}", serviceUid=c["serviceUid"])
            tick(results[k])

    res_df = pd.DataFrame([{**dict(zip(TRIP_KEY, k)), **v} for k, v in results.items()])
    return expected.merge(res_df, on=TRIP_KEY, how="left")


# ============================================================
# Excel export
# ============================================================
def to_excel(report: pd.DataFrame, summary: pd.DataFrame, by_diag: pd.DataFrame) -> bytes:
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    fills = {STATUS_OK: "E3F4E8", STATUS_FAIL: "FBE3E4", STATUS_NC: "EEEEEE"}
    out = report_table(report, for_export=True)
    issues = out[out["Status"] != STATUS_OK]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        issues.to_excel(xw, sheet_name="Issues", index=False)
        out.to_excel(xw, sheet_name="All results", index=False)
        by_diag.to_excel(xw, sheet_name="By diagram", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2448D6")
                cell.alignment = Alignment(vertical="center")
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
            headers = [c.value for c in ws[1]]
            if "Status" in headers:
                si = headers.index("Status")
                for row in ws.iter_rows(min_row=2):
                    f = fills.get(row[si].value)
                    if f:
                        for cell in row:
                            cell.fill = PatternFill("solid", fgColor=f)
    return buf.getvalue()


# ============================================================
# Report shaping
# ============================================================
def report_table(report: pd.DataFrame, for_export: bool = False) -> pd.DataFrame:
    r = report.sort_values(["diagram_date", "date", "exp_dep", "headcode"], na_position="last")
    out = pd.DataFrame({
        "Status": r["status"],
        "Diagram date": r["diagram_date"],
        "Diagram day": r["diagram_day"],
        "Runs on": r["date"],
        "Run day": r["run_day"] if "run_day" in r else r["diagram_day"],
        "After midnight": r["after_midnight"].map({True: "Y", False: ""}),
        "Headcode": r["headcode"],
        "From": r["exp_from"],
        "To": r["exp_to"],
        "Booked dep": r["exp_dep"],
        "RTT dep": r["act_dep"],
        "Booked arr": r["exp_arr"],
        "RTT arr": r["act_arr"],
        "Operator": r["operator"],
        "Issue": r["error"],
        "Depot": r["depot"],
        "Diagram": r["diagram"],
        "Resource": r["resource"],
        "RTT UID": r["serviceUid"],
    })
    if not for_export:
        out = compact_view(out)
    return out.reset_index(drop=True)


def _time_pair(booked, rtt) -> str:
    b = booked if isinstance(booked, str) else ""
    a = rtt if isinstance(rtt, str) else ""
    if b and a and a != b:
        return f"{b} ≠ {a}"
    return b or a or "—"


def compact_view(out: pd.DataFrame) -> pd.DataFrame:
    """On-screen table: few, narrow columns so it fits the page without sideways scrolling."""
    icon = {STATUS_OK: "✅ OK", STATUS_FAIL: "❌ FAIL", STATUS_NC: "⚪ NOT CHECKED"}
    short = lambda iso: pd.to_datetime(iso).strftime("%a %d %b")
    runs = [short(r) + (" 🌙" if m == "Y" else "") if m == "Y" else ""
            for r, m in zip(out["Runs on"], out["After midnight"])]
    return pd.DataFrame({
        "Status": out["Status"].map(icon),
        "Diagram": out["Diagram"],
        "Day": out["Diagram date"].map(short),
        "Headcode": out["Headcode"],
        "Journey": out["From"] + " → " + out["To"],
        "Dep": [_time_pair(b, a) for b, a in zip(out["Booked dep"], out["RTT dep"])],
        "Arr": [_time_pair(b, a) for b, a in zip(out["Booked arr"], out["RTT arr"])],
        "Issue": out["Issue"].fillna(""),
        "Runs on": runs,                       # only filled for after-midnight trains
        # used by the filters, not shown
        "Depot": out["Depot"], "Diagram day": out["Diagram day"], "Diagram date": out["Diagram date"],
        "Run day": out["Run day"],
    })


SHOWN_COLS = ["Status", "Day", "Headcode", "Journey", "Dep", "Arr", "Issue", "Runs on", "Diagram"]
COL_CFG = {
    "Status": st.column_config.TextColumn(width="small") if hasattr(st, "column_config") else None,
    "Day": st.column_config.TextColumn(width="small") if hasattr(st, "column_config") else None,
    "Headcode": st.column_config.TextColumn(width="small") if hasattr(st, "column_config") else None,
    "Dep": st.column_config.TextColumn("Dep (booked ≠ RTT)", width="small") if hasattr(st, "column_config") else None,
    "Arr": st.column_config.TextColumn("Arr (booked ≠ RTT)", width="small") if hasattr(st, "column_config") else None,
    "Issue": st.column_config.TextColumn(width="large") if hasattr(st, "column_config") else None,
    "Runs on": st.column_config.TextColumn("Runs on 🌙", width="small") if hasattr(st, "column_config") else None,
}


def show_table(df: pd.DataFrame, cols: List[str] = SHOWN_COLS, height: Optional[int] = None):
    cfg = {k: v for k, v in COL_CFG.items() if v is not None and k in cols}
    kw = {"height": height} if height else {}
    st.dataframe(style_table(df), use_container_width=True, hide_index=True,
                 column_order=cols, column_config=cfg, **kw)


def diagram_summary(report: pd.DataFrame) -> pd.DataFrame:
    g = report.groupby(["depot", "diagram_id", "plan_type", "diagram_days", "diagram"], dropna=False)
    s = g["status"].agg(
        Checks="size",
        OK=lambda x: (x == STATUS_OK).sum(),
        FAIL=lambda x: (x == STATUS_FAIL).sum(),
        NOT_CHECKED=lambda x: (x == STATUS_NC).sum(),
    ).reset_index()
    s = s.rename(columns={"depot": "Depot", "diagram_id": "Diagram ID", "plan_type": "Plan",
                          "diagram_days": "Days", "diagram": "Diagram", "NOT_CHECKED": "NOT CHECKED"})
    return s.sort_values(["FAIL", "NOT CHECKED", "Depot", "Diagram ID"], ascending=[False, False, True, True])


# ============================================================
# UI
# ============================================================
CSS = """
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Familjen+Grotesk:wght@500;600;700&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap">
<style>
/* Same palette and fonts as the Traincrew Planning Hub */
:root { --accent:#2448D6; --accent-soft:#E2E8FC; --navy:#0C1424; --ink:#152030;
  --ok:#167449; --ok-soft:#DFF2E7; --warn:#B42318; --warn-soft:#FCE6E3; --bh:#7339BF; --bh-soft:#EEE4FA;
  --f-display:"Familjen Grotesk","Helvetica Neue",Arial,sans-serif;
  --f-body:"IBM Plex Sans",system-ui,-apple-system,"Segoe UI",sans-serif;
  --f-mono:"IBM Plex Mono",ui-monospace,Menlo,monospace; }
html, body, [class*="css"], .stMarkdown, .stDataFrame, button, input, select, textarea { font-family: var(--f-body); }
h1, h2, h3, h4 { font-family: var(--f-display) !important; letter-spacing: -.01em; }
.block-container { padding-top: 1.6rem; max-width: 1500px; }
.prc-hero { background: linear-gradient(110deg, var(--navy) 0%, #1A2C5C 55%, var(--accent) 100%); color: #fff;
  border-radius: 14px; padding: 18px 24px; margin-bottom: 18px; }
.prc-hero h1 { color:#fff; font: 700 1.7rem/1.1 var(--f-display); margin: 0; }
.prc-hero p { margin: 4px 0 0; color: #DCE3F0; font-size: .92rem; }
.prc-tiles { display: grid; grid-template-columns: repeat(auto-fit, minmax(160px, 1fr)); gap: 12px; margin: 6px 0 16px; }
.prc-tile { border-radius: 12px; padding: 14px 16px; border: 1px solid rgba(90,101,119,.28); background: rgba(90,101,119,.06); }
.prc-tile .v { font: 700 1.9rem/1.1 var(--f-display); font-variant-numeric: tabular-nums; }
.prc-tile .l { font: 500 .72rem var(--f-mono); text-transform: uppercase; letter-spacing: .06em; opacity: .75; }
.prc-tile.ok   { border-color: rgba(22,116,73,.45);  background: rgba(22,116,73,.10); }
.prc-tile.ok .v { color: var(--ok); }
.prc-tile.fail { border-color: rgba(180,35,24,.45);  background: rgba(180,35,24,.10); }
.prc-tile.fail .v { color: var(--warn); }
.prc-tile.nc   { border-color: rgba(90,101,119,.45); }
.prc-tile.moon { border-color: rgba(115,57,191,.45); background: rgba(115,57,191,.10); }
.prc-tile.moon .v { color: var(--bh); }
section[data-testid="stSidebar"] h3 { font: 600 .78rem var(--f-mono) !important; text-transform: uppercase;
  letter-spacing: .08em; opacity: .8; margin-top: .8rem; }
.prc-wk { display: flex; gap: 14px; align-items: center; border: 1px solid rgba(128,128,128,.35);
  border-radius: 10px; padding: 10px 12px; }
.prc-wk .n { font: 700 1.9rem/1 var(--f-display); color: #4A6CF0; }
.prc-wk .d { font-size: .82rem; opacity: .8; line-height: 1.4; }
</style>
"""


def tiles(items: List[Tuple[str, object, str]]):
    html = "".join(f'<div class="prc-tile {cls}"><div class="v">{v}</div><div class="l">{l}</div></div>'
                   for l, v, cls in items)
    st.markdown(f'<div class="prc-tiles">{html}</div>', unsafe_allow_html=True)


def style_table(df: pd.DataFrame):
    def row_style(row):
        s = str(row["Status"])
        bg = ("rgba(180,35,24,.10)" if "FAIL" in s else
              "rgba(128,128,128,.12)" if "NOT" in s else "")
        styles = [f"background-color: {bg}" if bg else "" for _ in row]
        for c in ("Dep", "Arr"):
            if c in row.index and "≠" in str(row[c]):
                styles[list(row.index).index(c)] = "background-color: rgba(180,35,24,.30); font-weight: 700"
        return styles
    return df.style.apply(row_style, axis=1)


@st.cache_resource(show_spinner=False)
def get_live_client(user: str, pw: str) -> RTTClient:
    return RTTClient(user, pw)



@st.cache_data(show_spinner=False)
def cached_railrefs_from_repo(path: str) -> pd.DataFrame:
    return read_railrefs(path)


def filters(df: pd.DataFrame, prefix: str) -> pd.DataFrame:
    c1, c2, c3, c4 = st.columns([1.2, 1, 1.2, 1])
    status = c1.multiselect("Status", sorted(df["Status"].unique()), key=f"{prefix}_st")
    depot = c2.multiselect("Depot", sorted(x for x in df["Depot"].unique() if x), key=f"{prefix}_dp")
    present = set(df["Diagram day"]) | set(df["Run day"])
    day = c3.multiselect("Day", [d for d in RAIL_ORDER if d in present], key=f"{prefix}_dy",
                         help="Diagram day, plus after-midnight trains that run on that day")
    q = c4.text_input("Headcode / station", key=f"{prefix}_q").strip().lower()
    if status:
        df = df[df["Status"].isin(status)]
    if depot:
        df = df[df["Depot"].isin(depot)]
    if day:
        df = df[df["Diagram day"].isin(day) | df["Run day"].isin(day)]
    if q:
        df = df[df[["Headcode", "Journey", "Diagram"]].apply(lambda r: q in " ".join(map(str, r)).lower(), axis=1)]
    return df


def day_picker(key_: str) -> List[str]:
    """Saturday-to-Friday day picker (all days selected by default)."""
    if hasattr(st, "pills"):
        picked = st.pills("Diagram days", DAY_SHORT, selection_mode="multi", default=DAY_SHORT,
                          key=key_, label_visibility="collapsed")
    else:
        picked = st.multiselect("Diagram days", DAY_SHORT, default=DAY_SHORT, key=key_,
                                label_visibility="collapsed")
    return [d for d in DAY_SHORT if d in (picked or [])]


def main():
    st.set_page_config(page_title="Pass Ride Checker · LNER (TRENT)", page_icon="🚆", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown(
        '<div class="prc-hero"><h1>🚆 Pass Ride Checker</h1>'
        '<p>LNER (TRENT) · checks other operators\' trains in pass-trips.csv against Realtime Trains</p></div>',
        unsafe_allow_html=True,
    )

    try:
        rtt_user = st.secrets.get("RTT_USER", "")
        rtt_pass = st.secrets.get("RTT_PASS", "")
    except Exception:   # no secrets file at all
        rtt_user = rtt_pass = ""
    try:
        default_workers = int(st.secrets.get("RTT_WORKERS", DEFAULT_WORKERS))
    except Exception:
        default_workers = DEFAULT_WORKERS
    default_workers = max(1, min(MAX_WORKERS, default_workers))

    if not (rtt_user and rtt_pass):
        st.error("The Realtime Trains API login is missing. In Streamlit, open the app menu → "
                 "**Settings → Secrets** and add `RTT_USER` and `RTT_PASS`, then reload the app.")
        st.stop()

    # ---------------- Sidebar: files ----------------
    with st.sidebar:
        st.markdown("### 1 · Pass trips file")
        pass_file = st.file_uploader("pass-trips.csv", type=["csv"], label_visibility="collapsed")

        st.markdown("### 2 · RailReferences")
        refs_file = st.file_uploader("Upload an updated RailReferences.csv (optional)", type=["csv"], key="refs")
        refs_note = st.empty()
        refs_note.caption("Leave empty to use the RailReferences.csv saved with the app.")

    # ---------------- Load files ----------------
    try:
        rail_refs = read_railrefs(refs_file) if refs_file is not None else cached_railrefs_from_repo(DEFAULT_RAILREFS_PATH)
    except Exception as e:
        st.error(f"Couldn't read RailReferences.csv. Upload a copy in the sidebar. ({e})")
        st.stop()
    desc_to_crs = build_desc_to_crs(rail_refs)

    if not pass_file:
        st.info("👈 Upload **pass-trips.csv** in the sidebar. The railway week is picked up from the file.")
        st.stop()

    raw = load_pass_trips(pass_file)
    cols, missing = map_columns(raw)
    if missing:
        st.error("Missing required columns: " + ", ".join(missing))
        with st.expander("Columns found in the file"):
            st.write(list(raw.columns))
        st.stop()
    trips = prepare_trips(raw, cols, desc_to_crs)
    stations = set(trips["exp_from"]) | set(trips["exp_to"])
    n_missing = len(set(trips.loc[trips["origin_crs"].isna(), "exp_from"]) | set(trips.loc[trips["dest_crs"].isna(), "exp_to"]))
    which = "your uploaded copy" if refs_file is not None else "the app's saved copy"
    refs_note.caption(f"Using {which}. " + (f"All {len(stations)} stations matched." if not n_missing
                      else f"{len(stations) - n_missing} of {len(stations)} stations matched – see the warning on the right."))

    # ---------------- Sidebar: week, days, connection ----------------
    file_id = f"{pass_file.name}_{pass_file.size}"
    from_file = detect_rail_week(trips)
    detected = from_file or rail_week_of(dt.date.today())
    rail_year = detected[0]
    with st.sidebar:
        st.markdown("### 3 · Railway week")
        week = detected[1]
        week_dates = rail_week_dates(rail_year, week)
        source = "from the file" if from_file else "no diagram dates in the file – using this week"
        st.markdown(
            f'<div class="prc-wk"><span class="n">WK{week:02d}</span><span class="d">'
            f'Sat {week_dates[0]:%d %b} – Fri {week_dates[-1]:%d %b %Y}<br>'
            f'Rail year {rail_year}/{str(rail_year + 1)[-2:]} · {source}</span></div>',
            unsafe_allow_html=True)

        st.markdown("### 4 · Diagram days")
        days = day_picker(f"days_{file_id}")
        st.caption("Trains after midnight on a diagram are checked on the following day automatically.")

        st.markdown("### RTT")
        st.caption("Old Realtime Trains API – Realtime Trains switches it off on 31 Mar 2027. "
                   "Use app2.py (new API) from then on.")
        saved = remembered_workers()
        if saved is not None and not st.session_state.get("_workers_loaded"):
            st.session_state["workers"] = saved          # apply the browser's saved default once
            st.session_state["_workers_loaded"] = True
        st.session_state.setdefault("workers", default_workers)
        with st.expander(f"Advanced · {st.session_state['workers']} parallel requests"):
            workers = st.slider("Parallel RTT requests", 1, MAX_WORKERS, key="workers",
                                help="Higher is faster. The app slows down by itself if RTT asks it to.")
            if streamlit_js_eval is not None:
                if st.button(f"Save {workers} as my default", use_container_width=True):
                    streamlit_js_eval(js_expressions=f"localStorage.setItem('{WORKERS_STORE}', '{workers}')",
                                      key=f"save_workers_{workers}_{time.time()}")
                    st.toast(f"Saved. The app will start with {workers} parallel requests in this browser.")
                st.caption(f"Saved default in this browser: {saved if saved is not None else 'none'} "
                           f"(app default {default_workers}).")
            else:
                st.caption(f"Default {default_workers}. Set RTT_WORKERS in Secrets to change it.")

    wk_label = f"WK{week:02d}"
    out_name = f"{wk_label} PASS RIDES"
    days_label = "Sat–Fri" if len(days) == 7 else ", ".join(days) if days else "no days"
    heading = f"{wk_label} · {days_label} · {week_dates[0]:%d %b} – {week_dates[-1]:%d %b %Y}"

    if not days:
        st.warning("Pick at least one diagram day in the sidebar.")
        st.stop()

    expected = expected_for_days(trips, week_dates, days)

    fingerprint = (file_id, week, tuple(days), getattr(refs_file, "name", None), getattr(refs_file, "size", None))
    if st.session_state.get("fingerprint") != fingerprint:
        st.session_state.pop("report", None)       # inputs changed: old results are stale
        st.session_state["fingerprint"] = fingerprint

    n_unique = len(expected.drop_duplicates(TRIP_KEY)) if not expected.empty else 0
    unmapped = sorted(set(trips.loc[trips["origin_crs"].isna(), "exp_from"]) |
                      set(trips.loc[trips["dest_crs"].isna(), "exp_to"]))

    report = st.session_state.get("report")

    if report is None:
        st.markdown(f"#### Ready to check · {heading}")
        tiles([
            ("Brand-blank trips in file", len(trips), ""),
            ("Diagram rows to check", len(expected), ""),
            ("Trains to look up in RTT", n_unique, ""),
            ("After midnight", int(expected["after_midnight"].sum()) if not expected.empty else 0, "moon"),
        ])
        st.caption("One train can appear on several diagrams. Each train is looked up in RTT once, "
                   "and its result is copied to every diagram row it appears on.")
        if unmapped:
            st.warning("These stations have no CRS in RailReferences and won't be checked: " + ", ".join(unmapped))
        if expected.empty:
            st.warning("No Brand-blank trips run on the chosen diagram days.")
            st.stop()

        run = st.button(f"▶ Run check ({n_unique} trains)", type="primary", use_container_width=True)
        with st.expander("Trips that will be checked"):
            preview = expected[["diagram_date", "diagram_day", "date", "after_midnight", "headcode",
                                "exp_from", "exp_to", "exp_dep", "exp_arr", "diagram"]]
            st.dataframe(preview, use_container_width=True, hide_index=True)
        if not run:
            st.stop()

        client = get_live_client(rtt_user, rtt_pass)
        bar = st.progress(0.0, text="Starting…")
        t0 = time.time()

        def on_progress(done, total, fails):
            bar.progress(done / max(total, 1), text=f"Checked {done} of {total} trains · {fails} failed so far")

        report = run_checks(expected, client, workers=workers, progress=on_progress)
        bar.empty()
        st.session_state["report"] = report
        st.session_state["elapsed"] = time.time() - t0

    # ---------------- Results ----------------
    per_train = report.drop_duplicates(TRIP_KEY)
    t_ok = int((per_train["status"] == STATUS_OK).sum())
    t_fail = int((per_train["status"] == STATUS_FAIL).sum())
    t_nc = int((per_train["status"] == STATUS_NC).sum())
    r_bad = int((report["status"] != STATUS_OK).sum())
    elapsed = st.session_state.get("elapsed", 0)

    head_l, head_r = st.columns([4, 1])
    head_l.markdown(f"#### Results · {heading}")
    if head_r.button("↺ New check", use_container_width=True):
        st.session_state.pop("report", None)
        st.rerun()

    tiles([
        ("Trains OK", t_ok, "ok"),
        ("Trains failed", t_fail, "fail"),
        ("Trains not checked", t_nc, "nc"),
        ("After midnight", int(per_train["after_midnight"].sum()), "moon"),
        ("Checked in", f"{elapsed:.1f}s", ""),
    ])
    st.caption(f"{len(per_train)} trains, appearing on {len(report)} diagram rows. "
               f"The {t_fail + t_nc} trains with a problem appear on {r_bad} diagram rows.")
    if t_fail == 0 and t_nc == 0:
        st.success("All trains are running as diagrammed.")

    table = report_table(report)
    by_diag = diagram_summary(report)

    summary = build_summary(report, wk_label, week_dates, days_label)
    st.download_button(f"📊 Download {out_name}.xlsx", to_excel(report, summary, by_diag),
                       file_name=f"{out_name}.xlsx", type="primary", use_container_width=True,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       help="Sheets: Summary, Issues, All results, By diagram")

    t_issues, t_all, t_diag = st.tabs([f"⚠️ Issues ({r_bad} rows)", f"All results ({len(report)} rows)", "By diagram"])
    with t_issues:
        bad = table[~table["Status"].str.contains("OK")]
        if bad.empty:
            st.success("No issues found.")
        else:
            bad = filters(bad, "iss")
            show_table(bad, height=min(600, 38 + 35 * len(bad)))

    with t_all:
        f = filters(table, "all")
        show_table(f, height=600)

    with t_diag:
        st.dataframe(by_diag[["Diagram", "Checks", "OK", "FAIL", "NOT CHECKED"]], use_container_width=True,
                     hide_index=True, height=260)
        options = by_diag["Diagram"].tolist()
        labels = {d: (f"❌ {d}" if fl else ("⚪ " + d if nc else "✅ " + d))
                  for d, fl, nc in zip(by_diag["Diagram"], by_diag["FAIL"], by_diag["NOT CHECKED"])}
        sel = st.selectbox("Open a diagram", options, format_func=lambda d: labels[d])
        one = table[table["Diagram"] == sel]
        for ddate, grp in one.groupby("Diagram date", sort=True):
            st.markdown(f"**{grp['Diagram day'].iloc[0]} {pd.to_datetime(ddate):%d %b %Y}**")
            show_table(grp, ["Status", "Headcode", "Journey", "Dep", "Arr", "Issue", "Runs on"])



def build_summary(report: pd.DataFrame, wk_label: str, week_dates: List[dt.date], days_label: str) -> pd.DataFrame:
    per_train = report.drop_duplicates(TRIP_KEY)
    cnt = lambda df, s: int((df["status"] == s).sum())
    return pd.DataFrame({
        "Item": ["Railway week", "Diagram dates", "Diagram days",
                 "Trains checked", "Trains OK", "Trains failed", "Trains not checked",
                 "Diagram rows", "Diagram rows with an issue", "After-midnight trains", "Generated"],
        "Value": [wk_label, f"{week_dates[0]:%d/%m/%Y} to {week_dates[-1]:%d/%m/%Y}", days_label,
                  len(per_train), cnt(per_train, STATUS_OK), cnt(per_train, STATUS_FAIL), cnt(per_train, STATUS_NC),
                  len(report), int((report["status"] != STATUS_OK).sum()), int(per_train["after_midnight"].sum()),
                  dt.datetime.now().strftime("%d/%m/%Y %H:%M")],
    })


if __name__ == "__main__":
    main()
